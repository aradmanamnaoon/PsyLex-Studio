"""Conservative DOCX bilingual dictionary parser and cached JSONL ingester.

PDF-to-DOCX conversions can scramble columns and split terms over lines.
Suspect pairs are rejected rather than silently inventing terminology.
Coordinates are table/row or paragraph indices, never imagined page numbers.
"""
from __future__ import annotations
import json
import re
from pathlib import Path
from typing import Callable, Iterable
from docx import Document
from psylex.core.models import DictionaryEntry

LATIN = re.compile(r"[A-Za-z]")
PERSIAN = re.compile(r"[\u0600-\u06FF]")
PAIR = re.compile(r"(?P<english>[A-Za-z][A-Za-z\d\-\'()/. ,]{1,75}?)\s*[\t:–—]+\s*(?P<persian>[\u0600-\u06FF].{1,95})")
SPACES = re.compile(r"\s+")


def tidy(text: str) -> str:
    return SPACES.sub(" ", text.replace("\u200c", "‌").strip()).strip("|;:،. ")


def looks_english(text: str) -> bool:
    text = tidy(text)
    latin = len(LATIN.findall(text))
    farsi = len(PERSIAN.findall(text))
    return 2 <= latin <= 90 and farsi == 0 and len(text) <= 105 and latin / max(len(text), 1) > .45


def looks_persian(text: str) -> bool:
    text = tidy(text)
    fa = len(PERSIAN.findall(text))
    return 2 <= fa <= 125 and len(text) <= 170 and fa / max(len(text), 1) > .35 and not re.search(r"[A-Za-z]{3,}", text)


def candidate(english: str, persian: str, source: str, location: str,
              quality: float = 0.8) -> DictionaryEntry | None:
    en, fa = tidy(english), tidy(persian)
    if not looks_english(en) or not looks_persian(fa):
        return None
    if len(en.split()) > 10 or len(fa.split()) > 20:
        return None
    if en.lower() in {'english', 'translation', 'meaning', 'page'}:
        return None
    # Broken parentheses often indicate DOCX column corruption rather than a
    # trustworthy bilingual pair. Keep it out of the default translation index.
    if en.count('(') != en.count(')') or fa.count('(') != fa.count(')'):
        return None
    # Isolated digit fragments in this OCR conversion are mangled book text
    # (e.g. "0900 gestalt", "acculturati 2"), not trustworthy headwords.
    if re.search(r"(?:^\d+\s|\s\d+(?:\s|$))", en):
        return None
    return DictionaryEntry(en, fa, source, location, quality)


def parse_docx(path: Path, progress: Callable[[str], None] | None = None) -> list[DictionaryEntry]:
    """Extract provenance-anchored bilingual pairs without crossing DOCX columns."""
    doc = Document(path)
    label = path.name
    found: list[DictionaryEntry] = []
    seen: set[tuple[str, str, str]] = set()

    def add(en: str, fa: str, location: str, confidence: float) -> None:
        value = candidate(en, fa, label, location, confidence)
        if value is None:
            return
        key = (value.english.casefold(), value.persian, location)
        if key not in seen:
            found.append(value)
            seen.add(key)

    english_first = not path.name.startswith("source_b")
    for ti, table in enumerate(doc.tables):
        if progress and ti % 40 == 0:
            progress(f"Parsing {label}: table {ti + 1}/{len(doc.tables)}")
        for ri, row in enumerate(table.rows):
            cells = [tidy(cell.text) for cell in row.cells]
            loc = f"table {ti + 1}, row {ri + 1}"
            # Converted Word tables often contain independent bilingual columns.
            # In particular, a 4-cell row is [EN1, FA1, EN2, FA2]; sliding
            # 2-3 creates a *false* pair (EN2 -> FA1). Never use a sliding window.
            # We restrict extraction to physical pairs supported by the layout.
            if english_first:
                if len(cells) == 2:
                    pairs = [(0, 1)]
                elif len(cells) == 3:
                    # The middle and right cells normally form the right-hand pair;
                    # the left cell is unrelated text from the other book column.
                    pairs = [(1, 2)] + ([(0, 1)] if not cells[2] else [])
                elif len(cells) == 4:
                    pairs = [(0, 1), (2, 3)]
                elif len(cells) == 5:
                    # Malformed table conversions: each pair is commonly split
                    # across 1,3 (cell 2 blank) and 4,5, or 1,2 and 4,5.
                    pairs = ([(0, 2)] if not cells[1] else [(0, 1)]) + [(3, 4)]
                else:
                    pairs = []  # Unknown layout: prefer abstention to mistranslation.
            else:
                pairs = []  # Source B is paragraph-first, often definition tables.
            for a_index, b_index in pairs:
                a, b = cells[a_index], cells[b_index]
                loc_pair = f"{loc}, cells {a_index + 1}-{b_index + 1}"
                if looks_english(a) and looks_persian(b):
                    add(a, b, loc_pair, 0.94 if len(cells) <= 4 else 0.88)
                elif looks_persian(a) and looks_english(b):
                    add(b, a, loc_pair, 0.89 if len(cells) <= 4 else 0.85)
            for ci, cell in enumerate(cells):
                for m in PAIR.finditer(cell):
                    add(m['english'], m['persian'], loc + f", cell {ci + 1}", 0.72)
                # Some conversions preserve a single bilingual line with no tab.
                m = re.match(r"^([A-Za-z][A-Za-z\-()' ]{1,65})\s+([\u0600-\u06FF].+)$", cell)
                if m:
                    add(m[1], m[2], loc + f", cell {ci + 1}", 0.67)

    # A DOCX conversion sometimes splits a SINGLE dictionary entry across two
    # consecutive rows, duplicating English across its two source cells and
    # Persian across the next row's cells. Recover only structurally explicit
    # continuations in the SAME column pair; never look horizontally across
    # the two independent columns.
    if english_first:
        for ti, table in enumerate(doc.tables):
            rows = list(table.rows)
            for ri in range(len(rows) - 1):
                head = [tidy(cell.text) for cell in rows[ri].cells]
                tail = [tidy(cell.text) for cell in rows[ri + 1].cells]
                if len(head) != 4 or len(tail) != 4:
                    continue
                for first in (0, 2):
                    second = first + 1
                    english = head[first]
                    persian = tail[second]
                    if (english == head[second] and looks_english(english)
                            and looks_persian(persian)
                            and (tail[first] == persian or not tail[first])):
                        add(english, persian,
                            f"table {ti + 1}, rows {ri + 1}-{ri + 2}, cells {first + 1}-{second + 1}",
                            0.91)

    # Paragraphs can contain bilingual lines; don't pair unrelated adjacent
    # paragraphs because conversion can alternate independent columns.
    for pi, par in enumerate(doc.paragraphs):
        value = tidy(par.text)
        if not value:
            continue
        loc = f"paragraph {pi + 1}"
        if path.name.startswith("source_a"):
            continue  # The scanned English-first book has interleaved columns in paragraphs.
        for m in PAIR.finditer(value):
            add(m['english'], m['persian'], loc, 0.78)
        m = re.match(r"^([A-Za-z][A-Za-z\-()' ]{1,65})\s+([\u0600-\u06FF].+)$", value)
        if m:
            add(m[1], m[2], loc, 0.70)
        else:
            # Persian heading followed by an English equivalent in the same line.
            m = re.match(r"^([\u0600-\u06FF‌ ،\-]{2,65})\s+([A-Za-z][A-Za-z\-()' ]{2,65})$", value)
            if m:
                add(m[2], m[1], loc, 0.96)
    if progress:
        progress(f"Extracted {len(found)} candidate entries from {label}")
    return found


def write_jsonl(entries: Iterable[DictionaryEntry], path: Path) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open('w', encoding='utf-8') as out:
        for e in entries:
            out.write(json.dumps(e.to_dict(), ensure_ascii=False) + '\n')
            count += 1
    return count


def load_jsonl(path: Path) -> list[DictionaryEntry]:
    entries: list[DictionaryEntry] = []
    with path.open(encoding='utf-8') as src:
        for line in src:
            data = json.loads(line)
            data.pop('id', None)
            entries.append(DictionaryEntry(**data))
    return entries


def langchain_documents(path: Path) -> list:
    """Optional LangChain loader for paragraph-level non-tabular RAG.

    Table-aware structured extraction remains authoritative for exact terms.
    """
    from langchain_community.document_loaders import Docx2txtLoader
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    loaded = Docx2txtLoader(str(path)).load()
    splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=70)
    return splitter.split_documents(loaded)
