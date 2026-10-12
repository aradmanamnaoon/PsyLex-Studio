"""Create reproducible synthetic documents for release QA.

Usage: python -m scripts.make_document_fixtures --output samples/document_fixtures
Requires python-docx, PyMuPDF, Pillow (runtime document dependencies).
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from docx import Document
from PIL import Image, ImageDraw, ImageFont

TEXT = 'Cognitive behavioral therapy improves emotion regulation.'


def create(output: Path) -> Path:
    import pymupdf
    output.mkdir(parents=True, exist_ok=True)
    (output/'plain.txt').write_text(TEXT+'\n'+TEXT,encoding='utf-8')
    (output/'terms.csv').write_text('english\nAnxiety\nCognition\nAttention\n',encoding='utf-8')
    word=Document()
    word.add_heading('Clinical psychology',level=1)
    word.add_paragraph(TEXT)
    table=word.add_table(rows=2,cols=2)
    table.cell(0,0).text='Construct'; table.cell(0,1).text='Definition'
    table.cell(1,0).text='Working memory'; table.cell(1,1).text='Temporary storage and processing'
    word.save(output/'sample.docx')
    embedded=pymupdf.open()
    page=embedded.new_page();page.insert_text((60,95),TEXT,fontsize=17)
    embedded.save(output/'embedded_text.pdf');embedded.close()
    image=Image.new('RGB',(1800,350),'white')
    drawer=ImageDraw.Draw(image)
    options=[Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
             Path('C:/Windows/Fonts/arial.ttf')]
    available=next((p for p in options if p.exists()),None)
    font=ImageFont.truetype(str(available),55) if available else ImageFont.load_default()
    drawer.text((70,70),TEXT,fill='black',font=font)
    image.save(output/'scan.png')
    scanned=pymupdf.open()
    page=scanned.new_page(width=900,height=175)
    page.insert_image(page.rect,filename=str(output/'scan.png'))
    scanned.save(output/'scanned_text.pdf');scanned.close()
    manifest=[{'file':'scan.png','reference':TEXT},
              {'file':'scanned_text.pdf','reference':TEXT}]
    (output/'ocr_gold.jsonl').write_text('\n'.join(json.dumps(x) for x in manifest)+'\n',encoding='utf-8')
    print(f'Created 6 fixtures and OCR manifest in {output}')
    return output


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',default='samples/document_fixtures')
    args=parser.parse_args()
    create(Path(args.output))
