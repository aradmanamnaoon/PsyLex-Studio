"""Generate a clean sample bilingual DOCX for repeatable importer testing."""
from pathlib import Path
from docx import Document

SAMPLES = [
    ('cognitive psychology', 'روان‌شناسی شناختی'),
    ('working memory', 'حافظه کاری'),
    ('behavioral activation', 'فعال‌سازی رفتاری'),
    ('social anxiety', 'اضطراب اجتماعی'),
    ('attachment theory', 'نظریه دلبستگی'),
    ('executive function', 'کارکرد اجرایی'),
    ('metacognition', 'فراشناخت'),
    ('self-efficacy', 'خودکارآمدی'),
]


def main():
    path = Path(__file__).resolve().parents[1] / 'resources' / 'sample_dictionary.docx'
    doc = Document()
    doc.add_heading('Sample Psychology Terminology — EN ↔ FA', 0)
    doc.add_paragraph('Example data for import testing. Not clinically verified.')
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'English'
    table.rows[0].cells[1].text = 'Persian'
    for en, fa in SAMPLES:
        cells = table.add_row().cells
        cells[0].text = en
        cells[1].text = fa
    doc.save(path)
    print(path)


if __name__ == '__main__':
    main()
