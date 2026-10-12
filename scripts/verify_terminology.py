"""Exhaustive, repeatable offline PsyLex terminology audit.

Checks source provenance, extraction anomalies, dictionary ranking, lookup behavior,
and independent (human-maintained) reference cases. It does NOT assert that the
source dictionary's translations are semantically correct.

Usage: python scripts/verify_terminology.py --output reports/terminology_validation
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import tempfile
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from psylex.core.config import Settings
from psylex.core.database import Database, normalize
from psylex.core.ingest import load_jsonl
from psylex.rag.graph import TranslationPipeline

ARABIC = re.compile(r"[\u0600-\u06ff]")
LATIN = re.compile(r"[A-Za-z]")
WORD_LATIN = re.compile(r"[A-Za-z]{3,}")
TABLE_LOC = re.compile(r"^table (\d+), row (\d+)(?:, cells? (\d+)(?:-(\d+))?)?$")
PARA_LOC = re.compile(r"^paragraph (\d+)$")
CROSS_ROW_LOC = re.compile(r"^table (\d+), rows (\d+)-(\d+), cells (1-2|3-4)$")
SPACES = re.compile(r"\s+")


def plain(value: str) -> str:
    """Conservative Unicode normalization for finding source text, not editing text."""
    value = unicodedata.normalize("NFKC", value)
    value = value.translate(str.maketrans({'ي': 'ی', 'ك': 'ک', '\u200c': ' ', '\u200f': ' ', '\u200e': ' '}))
    return SPACES.sub(' ', value).casefold().strip()


def source_contains(haystack: str, needle: str) -> bool:
    """Compare whole normalized strings then contiguous text; avoid silent fuzzy matches."""
    haystack, needle = plain(haystack), plain(needle)
    return bool(needle) and needle in haystack


@dataclass(slots=True)
class SourceIndex:
    tables: list[list[list[str]]]
    paragraphs: list[str]
    name: str

    @classmethod
    def from_docx(cls, path: Path) -> 'SourceIndex':
        from docx import Document
        d = Document(path)
        return cls(
            tables=[[[cell.text for cell in row.cells] for row in table.rows] for table in d.tables],
            paragraphs=[paragraph.text for paragraph in d.paragraphs],
            name=path.name,
        )

    def locate(self, loc: str, english: str, persian: str) -> tuple[list[str], str]:
        """Return provenance flags and evidence strength; row co-occurrence alone is weak."""
        # A term broken over a physical row boundary is valid only when the
        # first row repeats the English headword across one pair of cells and
        # the next row carries Persian in those SAME cells.
        continuation = CROSS_ROW_LOC.fullmatch(loc)
        if continuation:
            ti, r1, r2 = (int(continuation[i]) - 1 for i in (1, 2, 3))
            start = 0 if continuation[4] == '1-2' else 2
            if not (0 <= ti < len(self.tables) and 0 <= r1 < r2 < len(self.tables[ti])
                    and r2 == r1 + 1):
                return ['source_location_out_of_bounds'], 'missing'
            head, tail = self.tables[ti][r1], self.tables[ti][r2]
            if len(head) != 4 or len(tail) != 4:
                return ['cited_cells_do_not_contain_pair'], 'missing'
            en_a, en_b = (plain(head[x]) for x in (start, start + 1))
            fa_a, fa_b = (plain(tail[x]) for x in (start, start + 1))
            if en_a == en_b and source_contains(head[start], english) and (
                    (fa_a == fa_b or not fa_a) and source_contains(tail[start + 1], persian)):
                return [], 'continuation_rows'
            return ['cited_cells_do_not_contain_pair'], 'unmatched'
        m = TABLE_LOC.fullmatch(loc)
        if m:
            ti, ri = int(m[1]) - 1, int(m[2]) - 1
            if ti < 0 or ti >= len(self.tables) or ri < 0 or ri >= len(self.tables[ti]):
                return ['source_location_out_of_bounds'], 'missing'
            cells = self.tables[ti][ri]
            e_pos = [i for i, c in enumerate(cells) if source_contains(c, english)]
            f_pos = [i for i, c in enumerate(cells) if source_contains(c, persian)]
            if not e_pos or not f_pos:
                flags = []
                if not e_pos:
                    flags.append('english_not_found_at_source_location')
                if not f_pos:
                    flags.append('persian_not_found_at_source_location')
                return flags, 'unmatched'
            # Validate the cells claimed by the ingestion parser, not merely the row.
            if m[3]:
                c1, c2 = int(m[3]) - 1, (int(m[4]) - 1 if m[4] else int(m[3]) - 1)
                if c1 < 0 or c2 >= len(cells):
                    return ['cited_cell_out_of_bounds'], 'unmatched'
                cited = [i for i in range(c1, c2 + 1)]
                if not any(i in cited for i in e_pos) or not any(i in cited for i in f_pos):
                    return ['cited_cells_do_not_contain_pair'], 'row_only'
                # In these 4-column dictionary tables, cells 1-2 and 3-4
                # are separate bilingual pairs. Sliding across cells 2-3
                # combines one entry's Persian with another's English.
                if self.name == 'source_a.docx' and len(cells) == 4 and (c1, c2) == (1, 2):
                    return ['crossed_independent_dictionary_columns'], 'cross_column'
            same = any(i in f_pos for i in e_pos)
            adjacent = any(abs(i - j) == 1 for i in e_pos for j in f_pos)
            if same:
                return [], 'same_cell'
            if adjacent:
                return [], 'adjacent_cells'
            # A 5-cell conversion can split the first bilingual pair across
            # cells 1 and 3, with an empty intermediate cell.
            if self.name == 'source_a.docx' and len(cells) == 5 and m[3]:
                if (int(m[3]), int(m[4] or m[3])) == (1, 3) and not plain(cells[1]):
                    return [], 'split_pair_across_empty_cell'
            return ['pair_not_adjacent_in_source_row'], 'row_only'
        m = PARA_LOC.fullmatch(loc)
        if m:
            p = int(m[1]) - 1
            if p < 0 or p >= len(self.paragraphs):
                return ['source_location_out_of_bounds'], 'missing'
            para = self.paragraphs[p]
            flags = []
            if not source_contains(para, english):
                flags.append('english_not_found_at_source_location')
            if not source_contains(para, persian):
                flags.append('persian_not_found_at_source_location')
            return flags, 'same_paragraph' if not flags else 'unmatched'
        return ['unrecognized_source_location'], 'missing'


def entry_flags(row: dict[str, Any]) -> list[str]:
    """Flags are review leads, not semantic judgments."""
    en, fa = str(row.get('english', '')), str(row.get('persian', ''))
    flags: list[str] = []
    if not en.strip() or not fa.strip():
        flags.append('empty_headword_or_translation')
    if not LATIN.search(en):
        flags.append('english_missing_latin')
    if not ARABIC.search(fa):
        flags.append('persian_missing_arabic_script')
    if ARABIC.search(en):
        flags.append('persian_characters_in_english')
    if WORD_LATIN.search(fa):
        flags.append('latin_word_in_persian')
    if re.search(r'\d', en) and re.search(r'[A-Za-z]\s+\d|\d\s+[A-Za-z]', en):
        flags.append('suspicious_digit_in_headword')
    if len(fa) > 95 or len(fa.split()) > 12:
        flags.append('unusually_long_persian_translation')
    if len(en) > 65 or len(en.split()) > 7:
        flags.append('unusually_long_english_headword')
    if en.count('(') != en.count(')') or fa.count('(') != fa.count(')'):
        flags.append('unbalanced_parentheses')
    if re.search(r'\uFFFD|[\x00-\x08\x0b\x0c\x0e-\x1f]', en + fa):
        flags.append('replacement_or_control_character')
    quality = row.get('quality')
    try:
        q = float(quality)
        if not 0 <= q <= 1:
            flags.append('quality_out_of_bounds')
        elif q < 0.75:
            flags.append('low_extraction_score')
    except (ValueError, TypeError):
        flags.append('invalid_extraction_score')
    return flags


CRITICAL = {
    'empty_headword_or_translation', 'english_missing_latin', 'persian_missing_arabic_script',
    'quality_out_of_bounds', 'invalid_extraction_score', 'source_location_out_of_bounds',
    'unrecognized_source_location', 'unknown_source_file', 'runtime_no_exact_translation',
    'runtime_selected_not_in_lexicon', 'runtime_mismatched_evidence', 'duplicate_exact_record',
}
HIGH = {
    'english_not_found_at_source_location', 'persian_not_found_at_source_location',
    'cited_cells_do_not_contain_pair', 'pair_not_adjacent_in_source_row',
    'crossed_independent_dictionary_columns',
    'latin_word_in_persian', 'persian_characters_in_english', 'replacement_or_control_character',
    'suspicious_digit_in_headword', 'unbalanced_parentheses',
}


def tier(flags: list[str]) -> str:
    return 'critical' if any(f in CRITICAL for f in flags) else ('high' if any(f in HIGH for f in flags) else ('review' if flags else 'none'))


def write_csv(path: Path, records: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as handle:
        w = csv.DictWriter(handle, fieldnames=columns, extrasaction='ignore')
        w.writeheader()
        for r in records:
            w.writerow({k: r.get(k, '') for k in columns})


def run_audit(
    lexicon: Path, sources: Path, gold: Path, output: Path, check_docx: bool = True,
) -> dict[str, Any]:
    """Audit all candidates and all unique headwords, writing reproducible reports."""
    output.mkdir(parents=True, exist_ok=True)
    with lexicon.open(encoding='utf-8') as file:
        rows = [json.loads(line) for line in file if line.strip()]
    by_word: dict[str, list[tuple[int, dict[str, Any]]]] = defaultdict(list)
    for i, row in enumerate(rows, 1):
        by_word[normalize(str(row.get('english', '')))].append((i, row))

    # Load each DOCX only once; never scan 600-page documents once per entry.
    known_docs: dict[str, SourceIndex] = {}
    if check_docx:
        for name in sorted({str(r.get('source', '')) for r in rows}):
            candidate_path = sources / name
            if name in {'source_a.docx', 'source_b.docx'} and candidate_path.is_file():
                known_docs[name] = SourceIndex.from_docx(candidate_path)

    reports: list[dict[str, Any]] = []
    by_flags: Counter[str] = Counter()
    evidence: Counter[str] = Counter()
    severity: Counter[str] = Counter()
    distinct_pairs: dict[str, set[str]] = {}
    for headword, matches in by_word.items():
        distinct_pairs[headword] = {plain(str(r['persian'])) for _, r in matches}
    for i, row in enumerate(rows, 1):
        reasons = entry_flags(row)
        en, fa = str(row.get('english', '')), str(row.get('persian', ''))
        name, location = str(row.get('source', '')), str(row.get('location', ''))
        status = 'unchecked'
        if check_docx:
            doc = known_docs.get(name)
            if doc is None:
                reasons.append('unknown_source_file')
                status = 'missing'
            else:
                more, status = doc.locate(location, en, fa)
                reasons.extend(more)
        evidence[status] += 1
        variants = distinct_pairs[normalize(en)]
        if len(variants) > 1:
            reasons.append('multiple_persian_equivalents')
        if sum(1 for _, e in by_word[normalize(en)] if plain(str(e['persian'])) == plain(fa)
               and e.get('source') == row.get('source') and e.get('location') == row.get('location')) > 1:
            reasons.append('duplicate_exact_record')
        level = tier(reasons)
        severity[level] += 1
        by_flags.update(set(reasons))
        reports.append(dict(
            row=i, english=en, persian=fa, source=name, location=location,
            extraction_score=row.get('quality', ''), source_alignment=status,
            persian_variants=len(variants), severity=level, flags=';'.join(sorted(set(reasons))),
        ))

    # Compare two source dictionaries without assuming one is ground truth.
    # Distinct valid Persian synonyms are common; discrepancies are review leads.
    cross_source: list[dict[str, str]] = []
    for word, grouped in sorted(by_word.items()):
        a = {plain(str(row['persian'])) for _, row in grouped if row['source'] == 'source_a.docx'}
        b = {plain(str(row['persian'])) for _, row in grouped if row['source'] == 'source_b.docx'}
        if a and b:
            shared = bool(a.intersection(b))
            cross_source.append(dict(english=str(grouped[0][1]['english']),
                source_a=' | '.join(sorted(a)), source_b=' | '.join(sorted(b)),
                agreement='yes' if shared else 'no'))
    write_csv(output / 'cross_source_comparison.csv', cross_source,
              ['english', 'source_a', 'source_b', 'agreement'])

    # Which candidates are actually selected by the offline application? This
    # matters more than counting every noisy alternative equally.
    selected_top_rows = []
    for word, grouped in sorted(by_word.items()):
        if not word:
            continue
        chosen_idx, _ = max(grouped, key=lambda pair: float(pair[1]['quality']))
        selected_top_rows.append(reports[chosen_idx - 1])
    top_high = sum(r['severity'] == 'high' for r in selected_top_rows)
    top_cross = sum('crossed_independent_dictionary_columns' in r['flags']
                    for r in selected_top_rows)
    alternative_available = sum(
        'crossed_independent_dictionary_columns' in reports[max(grouped, key=lambda pair: float(pair[1]['quality']))[0] - 1]['flags']
        and any('crossed_independent_dictionary_columns' not in reports[i - 1]['flags']
                for i, _ in grouped)
        for word, grouped in by_word.items() if word
    )
    write_csv(output / 'top_choice_risks.csv', sorted(
        (r for r in selected_top_rows if r['severity'] != 'none'),
        key=lambda r: ({'critical': 0, 'high': 1, 'review': 2}[r['severity']], r['row'])),
              ['row', 'english', 'persian', 'source', 'location', 'extraction_score',
               'source_alignment', 'persian_variants', 'severity', 'flags'])

    # Full execution path: import lexicon to temp SQLite and translate EVERY headword.
    runtime: list[dict[str, Any]] = []
    runtime_issues: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix='psylex_terminology_') as tmp:
        db = Database(Path(tmp) / 'terms.sqlite')
        inserted = db.insert_entries(load_jsonl(lexicon))
        translator = TranslationPipeline(db, Settings(mode='dictionary'))
        for headword, grouped in sorted(by_word.items()):
            if not headword:
                continue
            original = str(grouped[0][1]['english'])
            valid_targets = {str(r['persian']) for _, r in grouped}
            result = translator.translate(original, save=False)
            top_entry = db.exact(original, 1)
            errors = []
            if not result.target or not top_entry:
                errors.append('runtime_no_exact_translation')
            elif result.target not in valid_targets:
                errors.append('runtime_selected_not_in_lexicon')
            if top_entry and (result.target != top_entry[0].persian or not result.evidence
                              or result.evidence[0].source != top_entry[0].source):
                errors.append('runtime_mismatched_evidence')
            # Normalization behavior is exercised on every term; no history writes.
            case_variant = translator.translate('  ' + original.swapcase() + '  ', save=False)
            if case_variant.target != result.target:
                errors.append('runtime_case_or_whitespace_inconsistent')
            runtime.append(dict(english=original, actual=result.target,
                                source=result.evidence[0].source if result.evidence else '',
                                variant_count=len(distinct_pairs[headword]),
                                flags=';'.join(errors)))
            if errors:
                runtime_issues.append(runtime[-1])
                by_flags.update(errors)
        # Gold entries must be independently reviewed; they are not all 20k entries.
        references: list[dict[str, str]] = []
        if gold.is_file():
            with gold.open(encoding='utf-8-sig', newline='') as f:
                for g in csv.DictReader(f):
                    query, target = str(g.get('english', '')).strip(), str(g.get('persian', '')).strip()
                    if not query or not target:
                        continue
                    response = translator.translate(query, save=False)
                    references.append(dict(english=query, expected=target, actual=response.target,
                        match='yes' if plain(response.target) == plain(target) else 'no',
                        target_present_in_source='yes' if target in [r.persian for r in db.exact(query, 9999)] else 'no'))
        db.close()

    write_csv(output / 'all_entries.csv', reports,
              ['row','english','persian','source','location','extraction_score',
               'source_alignment','persian_variants','severity','flags'])
    write_csv(output / 'review_queue.csv', sorted((r for r in reports if r['flags']),
              key=lambda r: ({'critical': 0, 'high': 1, 'review': 2}[r['severity']], r['row'])),
              ['row','english','persian','source','location','extraction_score',
               'source_alignment','persian_variants','severity','flags'])
    write_csv(output / 'runtime_results.csv', runtime, ['english','actual','source','variant_count','flags'])
    write_csv(output / 'gold_comparison.csv', references,
              ['english','expected','actual','match','target_present_in_source'])
    critical_count = severity['critical']
    summary: dict[str, Any] = {
        'meaning': 'Static/source checks and runtime lookup tests, NOT independently verified semantic translation accuracy',
        'entries_checked': len(rows), 'unique_english_headwords_checked': len(runtime),
        'imported_entries_sqlite': inserted, 'source_provenance_enabled': check_docx,
        'source_alignment': dict(sorted(evidence.items())), 'entry_review_severity': dict(sorted(severity.items())),
        'entries_with_multiple_persian_variants': sum(len(distinct_pairs[normalize(r['english'])]) > 1 for r in rows),
        'headwords_with_multiple_persian_variants': sum(len(v) > 1 for v in distinct_pairs.values()),
        'rule_counts': dict(sorted(by_flags.items(), key=lambda item: (-item[1], item[0]))),
        'runtime_failures': len(runtime_issues),
        'selected_top_choices_high_risk': top_high,
        'selected_top_choices_cross_column': top_cross,
        'cross_column_top_choices_with_an_alternative': alternative_available,
        'headwords_in_both_source_dictionaries': len(cross_source),
        'shared_headwords_with_identical_equivalent': sum(r['agreement'] == 'yes' for r in cross_source),
        'gold_reference_cases': len(references),
        'gold_exact_matches': sum(r['match'] == 'yes' for r in references),
        'gold_case_details': references,
        'no_flags_count': severity['none'],
        'critical_flags_count': critical_count,
        'overall_semantic_correctness_verified': False,
    }
    (output / 'summary.json').write_text(json.dumps(summary, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    md = [
        '# PsyLex Studio — Exhaustive Terminology Audit', '',
        'This is an **automated integrity and lookup test**, not proof that Persian translations are linguistically correct.',
        '', '## Coverage', '',
        f"- Dictionary candidates examined: **{len(rows):,}** (all rows).",
        f"- Unique English headwords translated by the offline engine: **{len(runtime):,}** (all headwords).",
        f"- SQLite import: **{inserted:,}** records.",
        f"- Source-DOCX provenance check: **{'enabled' if check_docx else 'disabled'}**.",
        f"- Runtime failures: **{len(runtime_issues):,}**.",
        f"- Headwords whose actual top choice is high-priority for review: **{top_high:,}**.",
        f"- Headwords whose top choice crosses independent DOCX columns: **{top_cross:,}**.",
        f"- Of these, headwords with an alternative non-crossing source entry: **{alternative_available:,}**.",
        f"- Headwords shared by both DOCX dictionaries: **{len(cross_source):,}**.",
        f"- Shared headwords with an identical normalized Persian equivalent: **{summary['shared_headwords_with_identical_equivalent']:,}**.",
        f"- Source entries with no automated flags: **{severity['none']:,}**.",
        f"- Critical-issue records: **{critical_count:,}**.",
        f"- High-priority review records: **{severity['high']:,}**.",
        f"- Other review records: **{severity['review']:,}**.",
        f"- Headwords with multiple distinct Persian variants: **{summary['headwords_with_multiple_persian_variants']:,}**.",
        '', '## Source alignment', '',
        '| Evidence type | Entries |', '|---|---:|',
        *[f'| {name} | {count:,} |' for name, count in sorted(evidence.items())],
        '', '## Triggered rules', '', '| Rule | Records/queries |', '|---|---:|',
        *[f'| `{name}` | {count:,} |' for name, count in summary['rule_counts'].items()],
        '', '## Project reference examples', '',
        f"Exact matches: **{summary['gold_exact_matches']}/{len(references)}** (existing project CSV; not an externally certified gold standard).", '',
        '| English | Expected | Actual | Match |', '|---|---|---|---|',
        *[f"| {r['english']} | {r['expected']} | {r['actual']} | {r['match']} |" for r in references],
        '', '## Examples: potentially wrong translations selected by the app', '',
        '| English | Current Persian output | DOCX location |', '|---|---|---|',
        *[f"| {r['english']} | {r['persian']} | {r['source']}, {r['location']} |"
          for r in selected_top_rows if 'crossed_independent_dictionary_columns' in r['flags']][:12],
        '', '## Important interpretation', '',
        '- `multiple_persian_equivalents` is a **review candidate**, not inherently an error: synonyms and polysemy are legitimate.',
        '- `low_extraction_score` is the converter heuristic, **not translation accuracy**.',
        '- `source_alignment=adjacent_cells` confirms location/co-occurrence, but does not certify column-to-column semantic pairing.',
        '- `source_alignment=cross_column` indicates likely mispairing across independent side-by-side dictionary columns.',
        '- `source_alignment=row_only` cannot establish the right column relationship.',
        '- Reference CSV can contain questionable expectations and should be reviewed by a domain expert.',
        '- Nothing here validates unfamiliar English words omitted during DOCX extraction, or new paragraph translations from LLMs.',
        '- No source records are automatically changed or deleted.',
        '', '## Generated files', '',
        '- `all_entries.csv`: one row for every lexicon candidate.',
        '- `review_queue.csv`: flagged rows ordered by severity then original row.',
        '- `runtime_results.csv`: output for every unique English headword.',
        '- `gold_comparison.csv`: project benchmark cases.',
        '- `top_choice_risks.csv`: potentially incorrect translations currently preferred by the app.',
        '- `cross_source_comparison.csv`: where source A and B agree or disagree (synonyms may still be correct).',
        '- `summary.json`: complete machine-readable totals.',
        '', '## Follow-up', '',
        'Review high/critical rows in the DOCX, approve authoritative variants in a human-curated reference list,',
        'then rerun this script. Passing internal checks is necessary, **not sufficient**, for academic or clinical publication.',
    ]
    (output / 'REPORT.md').write_text('\n'.join(md) + '\n', encoding='utf-8')
    return summary


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--lexicon', type=Path, default=ROOT / 'resources' / 'starter_glossary.jsonl')
    ap.add_argument('--sources', type=Path, default=ROOT / 'resources' / 'dictionaries')
    ap.add_argument('--gold', type=Path, default=ROOT / 'tests' / 'gold.csv')
    ap.add_argument('--output', type=Path, default=ROOT / 'reports' / 'terminology_validation')
    ap.add_argument('--skip-docx', action='store_true', help='Skip provenance checks when source DOCX files are unavailable')
    ap.add_argument('--fail-on-runtime', action='store_true', help='Exit nonzero if any offline lookup behavior fails')
    ap.add_argument('--fail-on-quality', action='store_true', help='Exit nonzero if risky top choices or source column-crossing exist')
    ap.add_argument('--fail-on-gold', action='store_true', help='Exit nonzero if any reference case differs')
    args = ap.parse_args(argv)
    result = run_audit(args.lexicon, args.sources, args.gold, args.output, check_docx=not args.skip_docx)
    print(json.dumps({k: v for k, v in result.items() if k not in {'rule_counts','gold_case_details'}},
                     indent=2, ensure_ascii=False))
    print(f'Report: {args.output / "REPORT.md"}')
    if (args.fail_on_runtime and result['runtime_failures']) or (
        args.fail_on_gold and result['gold_exact_matches'] != result['gold_reference_cases']
    ) or (args.fail_on_quality and (result['selected_top_choices_high_risk'] or
                                    result['rule_counts'].get('crossed_independent_dictionary_columns', 0))):
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
