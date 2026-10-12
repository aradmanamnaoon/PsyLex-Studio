"""Regression suite for multi-format loading, translation, and output safety."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from docx import Document
from psylex.core.documents import (DocumentData, TextBlock, load_document,
                                   save_document, split_for_translation, translate_document)


class Response:
    def __init__(self, target: str):
        self.target, self.warnings = target, []


class FakeTranslator:
    def translate(self, text: str, save: bool = True):
        assert save is False
        return Response('ترجمه ' + text)


def test_word_load_and_export_with_table(tmp_path):
    source=tmp_path/'source.docx'
    word=Document()
    word.add_heading('Cognitive psychology', level=1)
    word.add_paragraph('Working memory and attention.')
    tab=word.add_table(rows=2,cols=2)
    tab.cell(0,0).text='Term'; tab.cell(0,1).text='Meaning'
    tab.cell(1,0).text='Anxiety'; tab.cell(1,1).text='Clinical construct'
    word.save(source)
    loaded=load_document(source)
    assert len(loaded.blocks)==6
    assert loaded.blocks[0].kind=='heading'
    assert loaded.blocks[2].kind=='table_cell'
    assert loaded.blocks[3].col==1
    translate_document(loaded, FakeTranslator())
    output=tmp_path/'persian.docx'
    save_document(loaded,output)
    check=Document(output)
    assert len(check.tables)==1
    assert check.tables[0].cell(1,0).text.startswith('ترجمه ')
    assert 'ترجمه' in check.paragraphs[1].text
    assert 'w:bidi' in check.paragraphs[1]._p.xml


def test_docx_original_cannot_be_overwritten(tmp_path):
    p=tmp_path/'same.docx'; Document().save(p)
    doc=DocumentData(str(p),'.docx',[TextBlock('Hello',translated='سلام')])
    with pytest.raises(ValueError,match='overwrite'):
        save_document(doc,p)


@pytest.mark.parametrize('extension', ['.txt','.md','.html','.json','.jsonl','.csv','.tsv'])
def test_round_trip_exports(tmp_path,extension):
    document=DocumentData(str(tmp_path/'src.txt'),'.txt',
        [TextBlock('Cognitive dissonance',translated='ناهماهنگی شناختی'),
         TextBlock('Working memory',translated='حافظه کاری')])
    result=tmp_path/('translated'+extension)
    save_document(document,result)
    text=result.read_text(encoding='utf-8-sig')
    assert 'ناهماهنگی شناختی' in text
    assert 'حافظه کاری' in text
    if extension in {'.json','.jsonl'}:
        parsed=json.loads(text) if extension=='.json' else [json.loads(line) for line in text.splitlines()]
        assert parsed[0]['source']=='Cognitive dissonance'


def test_csv_matrix_preserves_rows(tmp_path):
    src=tmp_path/'src.csv'
    src.write_text('Anxiety,Memory\nStress,Attention',encoding='utf-8')
    doc=load_document(src)
    translate_document(doc,FakeTranslator())
    dst=tmp_path/'out.csv'
    save_document(doc,dst)
    with dst.open(encoding='utf-8-sig',newline='') as fp:
        rows=list(csv.reader(fp))
    assert len(rows)==2 and all(len(row)==2 for row in rows)
    assert rows[0][0].startswith('ترجمه ')


def test_pdf_embedded_text(tmp_path):
    fitz=pytest.importorskip('fitz')
    path=tmp_path/'text.pdf'
    document=fitz.open(); page=document.new_page(); page.insert_text((50,80),'Cognitive behavioral therapy')
    document.save(path); document.close()
    loaded=load_document(path,ocr_mode='never')
    assert 'Cognitive behavioral therapy' in loaded.blocks[0].text
    assert not loaded.blocks[0].ocr
    assert loaded.blocks[0].page==1


def test_pdf_scanned_ocr(tmp_path):
    fitz=pytest.importorskip('fitz')
    from PIL import Image, ImageDraw, ImageFont
    import shutil
    if not shutil.which('tesseract'): pytest.skip('Native Tesseract executable unavailable')
    image=Image.new('RGB',(1350,240),'white')
    draw=ImageDraw.Draw(image)
    font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    font=ImageFont.truetype(font_path,55) if Path(font_path).exists() else ImageFont.load_default()
    draw.text((60,60),'Psychological assessment',font=font,fill='black')
    png=tmp_path/'image.png'; image.save(png)
    pdf=fitz.open(); page=pdf.new_page(width=675,height=120)
    page.insert_image(page.rect,filename=str(png)); path=tmp_path/'scanned.pdf'; pdf.save(path); pdf.close()
    loaded=load_document(path,ocr_mode='auto')
    assert loaded.blocks and loaded.blocks[0].ocr
    assert 'Psychological' in ' '.join(b.text for b in loaded.blocks)


def test_pdf_never_ocr_fails_for_scans(tmp_path):
    fitz=pytest.importorskip('fitz')
    pdf=fitz.open(); pdf.new_page(); file=tmp_path/'empty.pdf'; pdf.save(file); pdf.close()
    with pytest.raises(ValueError,match='No readable text'):
        load_document(file,ocr_mode='never')


def test_very_long_input_is_segmented_without_loss():
    text='Cognition memory attention. ' * 300
    chunks=split_for_translation(text,max_chars=180)
    assert all(len(c)<=180 for c in chunks)
    assert ''.join(''.join(c.split()) for c in chunks)==''.join(text.split())


def test_dictionary_only_fails_instead_of_exporting_blank(tmp_path):
    class Blank:
        def translate(self, text, save=False): return Response('')
    doc=DocumentData(str(tmp_path/'src.txt'),'.txt',[TextBlock('Complex clinical sentence')])
    with pytest.raises(RuntimeError,match='configure a local or API LLM'):
        translate_document(doc,Blank())
    with pytest.raises(ValueError,match='Translate all'):
        save_document(doc,tmp_path/'failed.docx')


def test_cancellation():
    doc=DocumentData('fixture.txt','.txt',[TextBlock('Hello')])
    with pytest.raises(InterruptedError):
        translate_document(doc, FakeTranslator(),is_cancelled=lambda:True)


def test_json_scalar_keys_are_not_renamed(tmp_path):
    path=tmp_path/'src.json'; path.write_text('{"term":"cognition","count":5,"is_active":true}',encoding='utf-8')
    loaded=load_document(path)
    assert len(loaded.blocks)==1 and loaded.blocks[0].text=='cognition'
    assert any('schema' in w for w in loaded.warnings)


def test_html_excludes_script(tmp_path):
    path=tmp_path/'a.html'; path.write_text('<p>memory</p><script>dangerous</script>',encoding='utf-8')
    loaded=load_document(path)
    assert 'memory' in loaded.blocks[0].text
    assert 'dangerous' not in loaded.blocks[0].text


def test_export_requires_all_segments(tmp_path):
    doc=DocumentData('original.txt','.txt',[TextBlock('Yes',translated='بله'), TextBlock('No')])
    with pytest.raises(ValueError): save_document(doc,tmp_path/'out.md')
    assert not (tmp_path/'out.md').exists()


def test_csv_header_is_preserved_in_dictionary_mode(tmp_path):
    src=tmp_path/'terms.csv'
    src.write_text('english\nAnxiety\nCognition',encoding='utf-8')
    doc=load_document(src)
    assert doc.blocks[0].kind=='table_header'
    translate_document(doc,FakeTranslator())
    assert doc.blocks[0].translated=='english'
    dest=tmp_path/'out.csv';save_document(doc,dest)
    assert dest.read_text(encoding='utf-8-sig').splitlines()[0]=='english'
