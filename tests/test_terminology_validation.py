"""Regression tests for the exhaustive terminology audit itself."""
import json
from pathlib import Path
from docx import Document
from scripts.verify_terminology import SourceIndex, entry_flags, run_audit


def test_source_alignment_and_wrong_pair(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for i, text in enumerate(['affect', 'عاطفه', 'attention', 'توجه']):
        table.cell(0, i).text = text
    for i, text in enumerate(['anxiety', 'اضطراب', '', '']):
        table.cell(1, i).text = text
    path = tmp_path / 'source_a.docx'
    doc.save(path)
    index = SourceIndex.from_docx(path)
    assert index.locate('table 1, row 1, cells 1-2', 'affect', 'عاطفه') == ([], 'adjacent_cells')
    assert index.locate('table 1, row 1, cells 2-3', 'attention', 'عاطفه') == (
        ['crossed_independent_dictionary_columns'], 'cross_column')
    assert index.locate('table 1, row 1, cells 1-2', 'affect', 'توجه')[0] == ['cited_cells_do_not_contain_pair']
    assert index.locate('table 1, row 1', 'affect', 'توجه')[0] == ['pair_not_adjacent_in_source_row']
    assert index.locate('table 1, row 3', 'affect', 'عاطفه')[0] == ['source_location_out_of_bounds']
    assert index.locate('paragraph 1', 'affect', 'عاطفه')[0] == ['source_location_out_of_bounds']


def test_entry_rules():
    assert not entry_flags({'english': 'working memory', 'persian': 'حافظه کاری', 'quality': 0.92})
    assert 'latin_word_in_persian' in entry_flags({'english': 'test', 'persian': 'روان test', 'quality': .9})
    assert 'low_extraction_score' in entry_flags({'english': 'test', 'persian': 'روان', 'quality': .65})
    assert 'persian_missing_arabic_script' in entry_flags({'english': 'test', 'persian': 'test', 'quality': .92})


def test_full_small_audit_and_correctness_benchmark(tmp_path):
    sources = tmp_path / 'sources'
    sources.mkdir()
    doc = Document()
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = 'cognitive therapy', 'درمان شناختی'
    doc.save(sources / 'source_a.docx')
    lexicon = tmp_path / 'starter_glossary.jsonl'
    lexicon.write_text(json.dumps({'english': 'cognitive therapy', 'persian': 'درمان شناختی',
                      'source': 'source_a.docx', 'location': 'table 1, row 1, cells 1-2',
                      'quality': .92}, ensure_ascii=False) + '\n', encoding='utf-8')
    gold = tmp_path / 'gold.csv'
    gold.write_text('english,persian\ncognitive therapy,درمان شناختی\n', encoding='utf-8')
    result = run_audit(lexicon, sources, gold, tmp_path / 'out')
    assert result['entries_checked'] == 1
    assert result['runtime_failures'] == 0
    assert result['gold_exact_matches'] == 1
    assert result['no_flags_count'] == 1
    assert result['selected_top_choices_cross_column'] == 0
    assert result['selected_top_choices_high_risk'] == 0
    assert (tmp_path / 'out' / 'all_entries.csv').is_file()
    assert (tmp_path / 'out' / 'REPORT.md').is_file()
