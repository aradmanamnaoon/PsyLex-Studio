from pathlib import Path
import json
import pytest
from docx import Document
from psylex.core.config import Settings
from psylex.core.database import Database, normalize
from psylex.core.dictionary import DictionaryService
from psylex.core.ingest import candidate, parse_docx, write_jsonl, load_jsonl
from psylex.core.models import DictionaryEntry, Feedback
from psylex.core.quality import qa_checks
from psylex.rag.graph import TranslationPipeline

@pytest.fixture
def db(tmp_path):
    handle = Database(tmp_path / 'test.db')
    handle.insert_entries([
        DictionaryEntry('working memory', 'حافظه کاری', 'unit.docx', 'table 1 row 2', .95),
        DictionaryEntry('working memory', 'حافظۀ فعال', 'unit2.docx', 'paragraph 5', .7),
        DictionaryEntry('cognitive therapy', 'درمان شناختی', 'unit.docx', 'table 1 row 3', .94),
        DictionaryEntry('attention', 'توجه', 'unit.docx', 'table 1 row 4', .95),
    ])
    yield handle
    handle.close()


def test_normalize():
    assert normalize('  WORKING   Memory ') == 'working memory'
    assert normalize("Child’s  Play") == "child's play"


def test_priority_and_provenance(db):
    result = TranslationPipeline(db, Settings()).translate('WORKING   MEMORY')
    assert result.target == 'حافظه کاری'
    assert result.evidence[0].source == 'unit.docx'
    assert result.confidence is not None


def test_correction_supersedes_dictionary(db):
    db.record_feedback(Feedback('working memory', 'حافظه کاری', False, 'حافظه فعال'))
    result = TranslationPipeline(db, Settings()).translate('working MEMORY')
    assert result.target == 'حافظه فعال'
    assert result.evidence[0].strategy == 'approved_correction'
    assert result.confidence == 1


def test_unknown_terms_do_not_hallucinate(db):
    hits = DictionaryService(db).search('completely unknown multiword term')
    assert not any(hit.kind in {'exact', 'approved'} for hit in hits)
    # Related suggestions must never be presented as verified dictionary definitions.
    assert all(hit.kind == 'related' for hit in hits)



def test_qa_preserves_numbers():
    assert any('numbers' in x for x in qa_checks('19 subjects', '۱۹ آزمودنی', [])) is False
    assert any('numbers' in x for x in qa_checks('19 subjects', '۲۰ آزمودنی', []))


def test_docx_sample_parsing(tmp_path):
    path = tmp_path / 'custom.docx'
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'English'
    table.cell(0, 1).text = 'Persian'
    table.cell(1, 0).text = 'attachment theory'
    table.cell(1, 1).text = 'نظریه دلبستگی'
    doc.save(path)
    entries = parse_docx(path)
    assert any(e.english == 'attachment theory' for e in entries)
    out = tmp_path / 'entries.jsonl'
    assert write_jsonl(entries, out) >= 1
    assert load_jsonl(out)[0].source == 'custom.docx'


def test_bad_pairs_are_rejected():
    assert candidate('hi', 'فارسی', 'x', 'x') is not None
    assert candidate('انگلیسی', 'English', 'x', 'x') is None
    assert candidate('a' * 500, 'فارسی', 'x', 'x') is None


def test_feedback_export_and_favorite(db, tmp_path):
    r = TranslationPipeline(db, Settings()).translate('attention')
    db.favorite(r.id, True)
    assert len(db.history(favorites_only=True)) == 1
    db.record_feedback(Feedback('attention', 'توجه', True))
    target = tmp_path / 'feedback.jsonl'
    assert db.export_feedback(target) == 1
    record = json.loads(target.read_text(encoding='utf-8'))
    assert record['chosen'] == 'توجه'
    assert record['rejected'] is None


def test_feedback_export_persian_utf8_roundtrip_on_any_os(db, tmp_path):
    """Windows ANSI code pages must never corrupt Persian review data."""
    db.record_feedback(Feedback('cognitive dissonance', 'ناسازگاری شناختی',
                                False, 'ناهماهنگی شناختی'))
    target = tmp_path / 'review.jsonl'
    assert db.export_feedback(target) == 1
    raw = target.read_bytes()
    assert 'ناهماهنگی شناختی'.encode('utf-8') in raw
    assert b'\xef\xbb\xbf' not in raw  # JSONL is plain UTF-8, no BOM.
    item = json.loads(raw.decode('utf-8'))
    assert item['chosen'] == 'ناهماهنگی شناختی'
    assert item['rejected'] == 'ناسازگاری شناختی'
    assert item['prompt'] == 'cognitive dissonance'


def test_feedback_export_survives_strict_unicode_decoding(db, tmp_path):
    """The exporter emits valid UTF-8 independent of OS locale."""
    db.record_feedback(Feedback('attention', 'توجه', True))
    target = tmp_path / 'feedback.jsonl'
    db.export_feedback(target)
    decoded = target.read_bytes().decode('utf-8', errors='strict')
    assert json.loads(decoded)['chosen'] == 'توجه'


def test_invalid_prompt_is_rejected():
    cfg = Settings(prompt='{text}')
    with pytest.raises(ValueError):
        cfg.validate()


def test_settings_roundtrip(tmp_path):
    path = tmp_path / 'settings.json'
    Settings(model='ollama/local-test', theme='light').save(path)
    loaded = Settings.load(path)
    assert loaded.theme == 'light'
    assert loaded.model == 'ollama/local-test'
    assert 'key' not in path.read_text(encoding='utf-8') or 'api_key_env' in path.read_text(encoding='utf-8')
