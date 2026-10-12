"""Safety regressions: never revive false DOCX pairings after future updates."""
from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document

from psylex.core.database import Database
from psylex.core.ingest import candidate, parse_docx
from psylex.core.models import DictionaryEntry, Feedback
from psylex.rag.graph import TranslationPipeline
from psylex.core.config import Settings


def test_four_columns_are_two_independent_pairs(tmp_path: Path) -> None:
    doc = Document()
    row = doc.add_table(rows=1, cols=4).rows[0]
    for i, text in enumerate(['abasement', 'خواری طلبی', 'attention', 'خطای مطلق']):
        row.cells[i].text = text
    source = tmp_path / 'source_a.docx'
    doc.save(source)
    terms = parse_docx(source)
    pairs = {(item.english, item.persian) for item in terms}
    assert pairs == {('abasement', 'خواری طلبی'), ('attention', 'خطای مطلق')}
    assert ('attention', 'خواری طلبی') not in pairs


def test_three_column_conversion_never_crosses_left_column(tmp_path: Path) -> None:
    doc = Document()
    row = doc.add_table(rows=1, cols=3).rows[0]
    for i, value in enumerate(['افسردگی', 'audibility range', 'دامنۀ شنودپذیری']):
        row.cells[i].text = value
    path = tmp_path / 'source_a.docx'
    doc.save(path)
    assert {(e.english, e.persian) for e in parse_docx(path)} == {
        ('audibility range', 'دامنۀ شنودپذیری')
    }


def test_five_column_split_pair_can_be_recovered_without_crossing(tmp_path: Path) -> None:
    doc = Document()
    row = doc.add_table(rows=1, cols=5).rows[0]
    for i, value in enumerate(['اسطوره ای', '', 'mythic', 'نزدیک بینی', 'myopia']):
        row.cells[i].text = value
    path = tmp_path / 'source_a.docx'
    doc.save(path)
    assert {(e.english, e.persian) for e in parse_docx(path)} == {
        ('mythic', 'اسطوره ای'), ('myopia', 'نزدیک بینی')
    }


def test_split_row_continuation_is_recovered(tmp_path: Path) -> None:
    from scripts.verify_terminology import SourceIndex
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for i, value in enumerate(['aberration', 'انحراف', 'absolute refractory period',
                                'absolute refractory period']):
        table.cell(0, i).text = value
    for i, value in enumerate(['abience', 'محرلگریزی', 'دورۀ بی پاسخی مطلق',
                                'دورۀ بی پاسخی مطلق']):
        table.cell(1, i).text = value
    path = tmp_path / 'source_a.docx'
    doc.save(path)
    entries = parse_docx(path)
    restored = next(e for e in entries if e.english == 'absolute refractory period')
    assert restored.persian == 'دورۀ بی پاسخی مطلق'
    assert restored.location == 'table 1, rows 1-2, cells 3-4'
    assert SourceIndex.from_docx(path).locate(restored.location, restored.english,
                                            restored.persian) == ([], 'continuation_rows')
    assert not any(e.english == 'absolute refractory period' and e.persian == 'انحراف'
                   for e in entries)


def test_parser_rejects_mixed_scripts_and_ocr_numbers() -> None:
    def ok(en: str, fa: str) -> bool:
        return candidate(en, fa, 'source_a.docx', 'table 1, row 1') is not None
    assert ok('attention', 'خطای مطلق')
    assert not ok('adjustment یل', 'سازگاری')
    assert not ok('aberration', 'تغییر affect')
    assert not ok('0900 gestalt', 'گشتالت خوب')
    assert not ok('abnormally (', 'نابهنجار')


def test_bundled_data_migration_preserves_user_work(tmp_path: Path) -> None:
    db = Database(tmp_path / 'migration.sqlite3')
    try:
        db.insert_entries([
            DictionaryEntry('attention', 'ناشناخته', 'PsyLex original starter glossary', 'starter entry 1', .78),
            DictionaryEntry('private term', 'ترجمه شخصی', 'personal.docx', 'table 1, row 1', .92),
        ])
        translation = TranslationPipeline(db, Settings()).translate('attention')
        db.favorite(translation.id, True)
        db.record_feedback(Feedback('private term', 'ترجمه شخصی', False, 'اصلاح شخصی'))
        corrected = [DictionaryEntry('attention', 'توجه', 'PsyLex original starter glossary', 'starter entry 1', .94)]
        fp = hashlib.sha256(b'v2').hexdigest()
        assert db.refresh_bundled_entries(corrected, fp) == 1
        assert db.refresh_bundled_entries(corrected, fp) == 0  # idempotence
        assert [(e.persian, e.source) for e in db.exact('attention')] == [
            ('توجه', 'PsyLex original starter glossary')
        ]
        assert db.exact('private term')[0].persian == 'ترجمه شخصی'
        assert db.correction('private term') == 'اصلاح شخصی'
        assert db.history(favorites_only=True)[0]['target'] == 'ناشناخته'
        assert db.count_entries() == 2
        assert TranslationPipeline(db, Settings()).translate('attention', save=False).target == 'توجه'
    finally:
        db.close()


def test_prompt_examples_require_a_related_english_token(tmp_path: Path) -> None:
    db = Database(tmp_path / 'examples.sqlite3')
    try:
        db.record_feedback(Feedback('working memory', 'اشتباه', False, 'حافظه کاری'))
        assert db.examples('working memory') == [('working memory', 'حافظه کاری')]
        assert db.examples('pyromania') == []  # no unrelated few-shot pollution
    finally:
        db.close()


def test_public_starter_glossary_is_searchable(tmp_path: Path) -> None:
    from psylex.core.ingest import load_jsonl
    source = Path(__file__).resolve().parents[1] / 'resources' / 'starter_glossary.jsonl'
    db = Database(tmp_path / 'public_terms.sqlite3')
    try:
        entries = load_jsonl(source)
        assert len(entries) >= 16
        assert all(e.source == 'PsyLex original starter glossary' for e in entries)
        db.insert_entries(entries)
        engine = TranslationPipeline(db, Settings(mode='dictionary'))
        assert engine.translate('attention', save=False).target == 'توجه'
        assert engine.translate('working memory', save=False).target == 'حافظه کاری'
    finally:
        db.close()
